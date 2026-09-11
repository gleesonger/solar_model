import csv
import sqlite3
from datetime import datetime
from collections import defaultdict
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook

fp_db: str = r"M:\media-server\config\solar-model\solar.db"

TIMEZONE: ZoneInfo = ZoneInfo("Europe/Dublin")
RESULTS_FILE = Path("tariff_model_results_ACTUALS.xlsx")

HOURLY_HEADERS = [
    "end_time_stamp", "local_time", "excel_timestamp", "supplier", "tariff_plan", "energy_plan",
    "period_sec", "load_wh", "solar_pv_wh", "battery_start_wh", "solar_to_load_wh",
    "solar_to_battery_wh", "grid_to_battery_wh", "battery_to_load_wh", "battery_to_grid_wh",
    "grid_import_wh", "grid_export_wh", "battery_end_wh", "import_c_per_kwh", "export_c_per_kwh",
    "import_cost_eur", "export_revenue_eur", "standing_charge_eur", "net_cost_eur",
]
DAILY_HEADERS = [
    "supplier", "tariff_plan", "date", "grid_import_kwh", "grid_export_kwh",
    "import_cost_eur", "export_revenue_eur", "standing_charge_eur", "net_cost_eur",
]
SUMMARY_HEADERS = DAILY_HEADERS[:2] + ["energy_plan", *DAILY_HEADERS[3:]]


def write_actuals_workbook(rows: list[list[object]]) -> None:
    """Write measured data using the same sheets and columns as model_tariffs."""
    rows = aggregate_hourly(rows)
    workbook = Workbook(write_only=True)
    hourly = workbook.create_sheet("hourly")
    hourly.append(HOURLY_HEADERS)
    for row in rows:
        hourly.append(row)

    daily: defaultdict[str, list[float]] = defaultdict(lambda: [0.0] * 6)
    for row in rows:
        day = str(row[1])[:10]
        values = daily[day]
        values[0] += float(row[15] or 0) / 1000
        values[1] += float(row[16] or 0) / 1000
        values[2] += float(row[20] or 0)
        values[3] += float(row[21] or 0)
        values[4] += float(row[22] or 0)
        values[5] += float(row[23] or 0)

    daily_sheet = workbook.create_sheet("daily_summary")
    daily_sheet.append(DAILY_HEADERS)
    for day, values in sorted(daily.items()):
        daily_sheet.append(["ACTUALS", "ACTUALS", day, *values])

    summary_sheet = workbook.create_sheet("summary")
    summary_sheet.append(SUMMARY_HEADERS)
    totals = [sum(values[index] for values in daily.values()) for index in range(6)]
    summary_sheet.append(["ACTUALS", "ACTUALS", "ACTUALS", *totals])
    workbook.save(RESULTS_FILE)


def aggregate_hourly(rows: list[list[object]]) -> list[list[object]]:
    grouped: dict[datetime, list[list[object]]] = defaultdict(list)
    for row in rows:
        local = datetime.fromisoformat(str(row[1]))
        grouped[local.replace(minute=0, second=0, microsecond=0)].append(row)

    result: list[list[object]] = []
    sum_columns = (6, 7, 8, 15, 16, 20, 21, 22, 23)
    for hour, hour_rows in sorted(grouped.items()):
        output = list(hour_rows[0])
        output[0] = int(hour.timestamp())
        output[1] = hour.isoformat()
        output[2] = hour.replace(tzinfo=None)
        for column in sum_columns:
            output[column] = sum(float(row[column] or 0) for row in hour_rows)
        output[9] = hour_rows[0][9]
        output[17] = hour_rows[-1][17]
        for column in (18, 19):
            rates = [float(row[column]) for row in hour_rows if row[column] is not None]
            output[column] = sum(rates) / len(rates) if rates else None
        result.append(output)
    return result

actual_rows: list[list[object]] = []
previous_battery_end_wh: float | None = None
with sqlite3.connect(fp_db) as db, open("minute_data.csv", "w", newline="") as output:
    writer = csv.writer(output)
    writer.writerow(["end_time_stamp", "excel_timestamp", "period_covered_sec", "load_wh", "solar_pv_wh"])

    previous_timestamp = None
    for row in db.execute(
        """
        SELECT
            collected_at_utc,
            COALESCE(plant_load_total_kwh_period, 0),
            COALESCE(plant_pv_total_kwh_period, 0),
            COALESCE(plant_grid_import_total_kwh_period, 0),
            COALESCE(plant_grid_export_total_kwh_period, 0),
            import_rate, export_rate,
            inverter_battery_available_discharge_kwh
        FROM sigenstor
        ORDER BY collected_at_utc
        """
    ):
        timestamp = row[0]
        load_period_raw = row[1]
        solar_period_raw = row[2]
        period = 0 if previous_timestamp is None else timestamp - previous_timestamp
        local = datetime.fromtimestamp(timestamp, TIMEZONE)
        load_wh = load_period_raw / 10
        solar_wh = solar_period_raw / 10
        grid_import_wh = row[3] / 10
        grid_export_wh = row[4] / 10
        # These columns are stored as fixed-point integers in the database;
        # the model workbook uses cents/kWh for rates and euros for costs.
        import_rate = row[5] / 100 if row[5] is not None else None
        export_rate = row[6] / 100 if row[6] is not None else None
        # Calculate costs from the normalized energy and rate units. Energy
        # is exported as Wh, rates as cents/kWh, and costs as euros.
        import_cost = (
            grid_import_wh / 1000 * import_rate / 100
            if import_rate is not None else None
        )
        export_revenue = (
            grid_export_wh / 1000 * export_rate / 100
            if export_rate is not None else None
        )
        net_cost = (
            import_cost - export_revenue
            if import_cost is not None and export_revenue is not None else None
        )
        # Stored kWh values use the database's fixed-point scale (10,000).
        # Convert the measured available battery level directly to Wh.
        battery_end_wh = row[7] / 10 if row[7] is not None else None
        battery_start_wh = previous_battery_end_wh if previous_battery_end_wh is not None else battery_end_wh
        writer.writerow([
            timestamp,
            local.replace(tzinfo=None).isoformat(sep=" "),
            period,
            load_wh,
            solar_wh,
        ])
        actual_rows.append([
            timestamp, local.isoformat(), local.replace(tzinfo=None), "ACTUALS", "ACTUALS", "ACTUALS",
            period, load_wh, solar_wh, battery_start_wh, None, None, None, None, None,
            grid_import_wh, grid_export_wh, battery_end_wh,
            import_rate, export_rate, import_cost, export_revenue, 0, net_cost,
        ])
        previous_battery_end_wh = battery_end_wh
        previous_timestamp = timestamp

write_actuals_workbook(actual_rows)
print(f"Writing {RESULTS_FILE}")
